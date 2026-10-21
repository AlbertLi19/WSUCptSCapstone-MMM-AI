"""Supervised pop-in detection trained on paired Excel depth/load annotations.

The desktop controls are in the application's existing Pop-in Locator tab.
"""
from __future__ import annotations

import argparse
import csv
import json
import hashlib
import os
import sys
from pathlib import Path

import joblib
import numpy as np
from sklearn.ensemble import RandomForestClassifier
from sklearn.model_selection import train_test_split
from scipy.ndimage import median_filter

ROOT = (Path(sys._MEIPASS) if getattr(sys, "frozen", False)
        else Path(__file__).resolve().parents[3])
DEFAULT_DATA = ROOT / "data" / "Pop-in Training Data"
DEFAULT_LABELS = ROOT / "data" / "pop-in 4H.xlsx"
BUNDLED_MODEL = ROOT / "models" / "popin_model.joblib"
DEFAULT_MODEL = ((Path(os.getenv("LOCALAPPDATA", str(Path.home()))) /
                  "SegmentationApp" / "models" / "popin_model.joblib")
                 if getattr(sys, "frozen", False) else BUNDLED_MODEL)
SHEET_FOLDERS = {
    "Off": "SiC-4H-OFF", "415 nm": "SiC-4H-415nm",
    "255 nm": "SiC-4H-UV255nm", "470 nm": "SiC-4H-UV470nm",
}
MODEL_VERSION = 2
FEATURE_SCHEMA = "normalized_local_shape_v2"


class DataError(ValueError):
    """Invalid curve, annotations, or model configuration."""


def load_curve(path):
    """Read physical depth (nm), load (uN), time (s), preserving row order.

    Instrument exports use a legacy micro symbol; decoding it must not damage
    numeric measurements. Voltage columns are deliberately not used.
    """
    path = Path(path)
    if path.suffix.lower() != ".txt" or not path.is_file():
        raise DataError("Select an existing .txt curve file.")
    text = path.read_bytes().decode("utf-8-sig", errors="replace")
    lines = text.splitlines()
    header = None
    for i, line in enumerate(lines):
        cols = [c.strip().lower() for c in line.split("\t")]
        if len(cols) >= 3 and cols[0] == "depth (nm)" and cols[1].startswith("load (") and cols[2] == "time (s)":
            unit_header = cols[1].replace("µ", "u").replace("μ", "u").replace("�", "u")
            if unit_header != "load (un)":
                raise DataError("Load must be in microNewtons (uN).")
            header = i
            break
    if header is None:
        raise DataError("Expected tab-separated Depth (nm), Load (uN), Time (s) columns.")
    rows = []
    for i, line in enumerate(lines[header + 1:], header + 2):
        if not line.strip():
            continue
        values = line.split("\t")
        if len(values) != len(cols):
            raise DataError(f"Line {i}: expected {len(cols)} columns.")
        try:
            row = [float(v) for v in values]
        except ValueError as exc:
            raise DataError(f"Line {i}: missing or nonnumeric measurement.") from exc
        if not np.isfinite(row).all():
            raise DataError(f"Line {i}: non-finite measurement.")
        rows.append(row[:3])
    curve = np.asarray(rows, dtype=float)
    if len(curve) < 50:
        raise DataError("At least 50 measurements are required.")
    if np.any(np.diff(curve[:, 2]) <= 0):
        raise DataError("Time must be strictly increasing.")
    if np.ptp(curve[:, 0]) == 0 or np.ptp(curve[:, 1]) == 0:
        raise DataError("Depth and load must both vary.")
    loading(curve)  # Validate loading portion before accepting the file.
    return curve


def loading(curve):
    end = int(np.argmax(curve[:, 1])) + 1
    if end < 40:
        raise DataError("At least 40 points before maximum load are required.")
    return curve[:end]


def features(curve):
    """Dimensionless local shape, invariant to axis scaling and offsets.

    No absolute coordinates or curve position are supplied to the forest.
    Forward changes are available because prediction uses a complete curve.
    """
    c = loading(curve)
    depth, force, time = c.T
    n = len(c)
    index = np.arange(n)
    dd, df = np.diff(depth), np.diff(force)
    depth_floor = np.ptp(depth) * 1e-9
    force_floor = np.ptp(force) * 1e-9
    baseline = np.median(dd)
    depth_step = max(np.median(np.abs(dd)), depth_floor)
    force_step = max(np.median(np.abs(df)), force_floor)
    noise = max(1.4826 * np.median(np.abs(dd - baseline)), depth_floor)
    increments = np.r_[dd[0], dd]
    local_baseline = median_filter(increments, size=41, mode="nearest")
    local_noise = np.maximum(1.4826 * median_filter(
        np.abs(increments - local_baseline), size=41, mode="nearest"), noise * 0.1)
    time_step = np.median(np.diff(time))
    columns = []
    for window in (1, 3, 5, 10, 20):
        before = np.maximum(index - window, 0)
        after = np.minimum(index + window, n - 1)
        for start, stop in ((before, index), (index, after)):
            dh, dl = depth[stop] - depth[start], force[stop] - force[start]
            steps = np.maximum(stop - start, 1)
            scaled_depth = dh / depth_step
            scaled_load = dl / force_step
            columns.extend([
                scaled_depth / steps,
                (dh - baseline * steps) / (noise * np.sqrt(steps)),
                scaled_load / steps,
                scaled_depth / np.maximum(np.abs(scaled_load), 0.05 * steps),
                (dh - local_baseline * steps) / (local_noise * np.sqrt(steps)),
                (time[stop] - time[start]) / (time_step * steps),
            ])
    return np.column_stack(columns).astype(np.float32)


def read_labels(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    labels = []
    try:
        for sheet, folder in SHEET_FOLDERS.items():
            if sheet not in workbook.sheetnames:
                raise DataError(f"Missing worksheet: {sheet}")
            rows = workbook[sheet].iter_rows(values_only=True)
            header = next(rows)
            if header[:2] != ("Depth", "Force (uN)"):
                raise DataError(f"{sheet}: expected Depth and Force (uN) in columns A/B.")
            # The contiguous table ends at the first blank A/B row. Numeric
            # means and standard deviations below it are not annotations.
            for row_number, row in enumerate(rows, 2):
                if row[0] is None and row[1] is None:
                    break
                try:
                    depth, force = float(row[0]), float(row[1])
                except (ValueError, TypeError) as exc:
                    raise DataError(f"{sheet}!A{row_number}: invalid label pair.") from exc
                if not np.isfinite([depth, force]).all():
                    raise DataError(f"{sheet}!A{row_number}: non-finite label.")
                labels.append(dict(sheet=sheet, folder=folder, excel_row=row_number,
                                   label_depth_nm=depth, label_load_uN=force))
    finally:
        workbook.close()
    if not labels:
        raise DataError("No pop-in annotations found.")
    return labels


def collect_training_data(data_dir, workbook, depth_tolerance=0.001, load_tolerance=0.01):
    """Scan all text files; use valid, annotated curves once each."""
    if depth_tolerance <= 0 or load_tolerance <= 0:
        raise DataError("Matching tolerances must be positive.")
    data_dir = Path(data_dir)
    if not data_dir.is_dir():
        raise DataError(f"Missing training directory: {data_dir}")
    labels = read_labels(workbook)
    matches = [[] for _ in labels]
    ambiguous = set()
    records, file_audit, fingerprints = {}, [], {}
    paths = [p for p in data_dir.rglob("*") if p.is_file() and p.suffix.lower() == ".txt"]
    # Prefer original numbered exports over duplicate alternate exports, but
    # accept arbitrary filenames. This keeps existing split membership stable.
    for path in sorted(paths, key=lambda p: (not p.stem[-4:].isdigit(), str(p))):
        relative = path.relative_to(data_dir)
        row = dict(file=str(relative), status="", reason="", matched_annotations=0)
        file_audit.append(row)
        try:
            curve = load_curve(path)
        except (DataError, OSError, UnicodeError) as error:
            row.update(status="invalid_format", reason=str(error))
            continue
        # Known treatment folders retain their sheet mapping; files elsewhere
        # in the training directory can match any worksheet by coordinates.
        folder = next((part for part in relative.parts[:-1] if part in SHEET_FOLDERS.values()), None)
        c = loading(curve)
        found = []
        for label_id, label in enumerate(labels):
            if folder is not None and label["folder"] != folder:
                continue
            indices = np.flatnonzero((np.abs(c[:, 0] - label["label_depth_nm"]) <= depth_tolerance)
                                     & (np.abs(c[:, 1] - label["label_load_uN"]) <= load_tolerance))
            if len(indices) == 1:
                found.append((label_id, int(indices[0])))
            elif len(indices) > 1:
                ambiguous.add(label_id)
        if not found:
            row.update(status="no_matching_annotation", reason="No unique Excel depth/load match in the loading segment.")
            continue
        sheets = {labels[label_id]["sheet"] for label_id, _ in found}
        if len(sheets) != 1:
            row.update(status="ambiguous_match", reason="Matches annotations from multiple treatment worksheets.")
            continue
        fingerprint = hashlib.sha256(curve.tobytes()).hexdigest()
        if fingerprint in fingerprints:
            row.update(status="duplicate_curve", reason=f"Identical measurements to {fingerprints[fingerprint]}.")
            continue
        fingerprints[fingerprint] = str(relative)
        records[path] = dict(curve=curve, indices=sorted({i for _, i in found}), sheet=next(iter(sheets)))
        row.update(status="matched", matched_annotations=len(found))
        for label_id, index in found:
            matches[label_id].append((str(relative), index,
                float(c[index, 0] - labels[label_id]["label_depth_nm"]),
                float(c[index, 1] - labels[label_id]["label_load_uN"])))
    audit = []
    for label_id, label in enumerate(labels):
        found = matches[label_id]
        audit.append(dict(label, status="matched" if found else ("ambiguous" if label_id in ambiguous else "unmatched"),
            file="; ".join(item[0] for item in found),
            index="; ".join(str(item[1]) for item in found),
            depth_error_nm="; ".join(str(item[2]) for item in found),
            load_error_uN="; ".join(str(item[3]) for item in found)))
    return records, audit, file_audit


def match_labels(data_dir, workbook, depth_tolerance=0.001, load_tolerance=0.01):
    records, audit, _ = collect_training_data(data_dir, workbook, depth_tolerance, load_tolerance)
    return records, audit


def event_indices(scores, threshold, separation=10):
    """Collapse adjacent detections, then suppress nearby duplicate events."""
    above = np.flatnonzero(scores >= threshold)
    if not len(above):
        return []
    runs = np.split(above, np.flatnonzero(np.diff(above) > 1) + 1)
    peaks = [int(run[np.argmax(scores[run])]) for run in runs]
    selected = []
    for peak in sorted(peaks, key=lambda i: (-scores[i], i)):
        if all(abs(peak - existing) > separation for existing in selected):
            selected.append(peak)
    return sorted(selected)


def evaluate(records, scores, threshold, tolerance=5):
    tp = fp = fn = 0
    errors = []
    for record, probabilities in zip(records, scores):
        truth = set(record["indices"])
        for i in sorted(event_indices(probabilities, threshold), key=lambda j: -probabilities[j]):
            available = [j for j in truth if abs(i - j) <= tolerance]
            if available:
                j = min(available, key=lambda j: abs(i - j))
                truth.remove(j)
                tp += 1
                errors.append(abs(i - j))
            else:
                fp += 1
        fn += len(truth)
    precision = tp / max(tp + fp, 1)
    recall = tp / max(tp + fn, 1)
    return dict(true_positives=tp, false_positives=fp, false_negatives=fn,
                precision=precision, recall=recall,
                f1=2 * precision * recall / max(precision + recall, 1e-12),
                mean_index_error=float(np.mean(errors)) if errors else None,
                location_tolerance_points=tolerance)


def train(data_dir=DEFAULT_DATA, workbook=DEFAULT_LABELS, model_path=DEFAULT_MODEL):
    data_dir, model_path = Path(data_dir), Path(model_path)
    records, audit, file_audit = collect_training_data(data_dir, workbook)
    model_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path = model_path.with_suffix(".matches.csv")
    with audit_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(audit[0]))
        writer.writeheader()
        writer.writerows(audit)
    with model_path.with_suffix(".files.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=["file", "status", "reason", "matched_annotations"])
        writer.writeheader()
        writer.writerows(file_audit)
    paths = sorted(records)
    if len(paths) < 20:
        raise DataError("At least 20 uniquely labeled curves are required for train/calibration/test splits.")
    groups = [records[p]["sheet"] for p in paths]
    try:
        development, test_paths = train_test_split(paths, test_size=0.2, random_state=42, stratify=groups)
        train_paths, calibration = train_test_split(
            development, test_size=0.25, random_state=43,
            stratify=[records[p]["sheet"] for p in development])
    except ValueError as error:
        raise DataError("Too few matched curves in one or more treatments for stratified validation. "
                        "Add labeled curves for those treatments; the existing model is preserved.") from error
    x, y = [], []
    for path in train_paths:
        record = records[path]
        f = features(record["curve"])
        target = np.zeros(len(f), dtype=np.uint8)
        for index in record["indices"]:
            target[max(0, index - 2):index + 3] = 1
        x.append(f)
        y.append(target)
    classifier = RandomForestClassifier(n_estimators=160, max_depth=16,
        min_samples_leaf=2, class_weight="balanced_subsample", random_state=42, n_jobs=-1)
    classifier.fit(np.vstack(x), np.concatenate(y))
    calibration_records = [records[p] for p in calibration]
    calibration_scores = [classifier.predict_proba(features(r["curve"]))[:, 1] for r in calibration_records]
    thresholds = np.linspace(0.05, 0.95, 37)
    calibration_metrics = [(float(t), evaluate(calibration_records, calibration_scores, t))
                           for t in thresholds]
    # Choose using calibration curves only. Equivalent F1/recall prefers the
    # lower cutoff rather than the most restrictive cutoff in a plateau.
    threshold = max(calibration_metrics, key=lambda item: (
        item[1]["f1"], item[1]["recall"], -item[0]))[0]
    test_records = [records[p] for p in test_paths]
    test_scores = [classifier.predict_proba(features(r["curve"]))[:, 1] for r in test_records]
    report = dict(model_version=MODEL_VERSION, feature_schema=FEATURE_SCHEMA,
        feature_count=classifier.n_features_in_,
        threshold_selection="calibration event F1, then recall, then lowest equivalent cutoff",
        annotation_count=sum(row["status"] == "matched" for row in audit),
        excel_annotation_count=len(audit),
        unmatched_annotation_count=sum(row["status"] != "matched" for row in audit),
        scanned_files=len(file_audit), skipped_files=[row for row in file_audit if row["status"] != "matched"],
        labeled_curves=len(paths), threshold=threshold,
        matching_tolerance=dict(depth_nm=0.001, load_uN=0.01),
        training_positive_radius_points=2, event_separation_points=10,
        split_files={"train": [str(p.relative_to(data_dir)) for p in train_paths],
                     "calibration": [str(p.relative_to(data_dir)) for p in calibration],
                     "test": [str(p.relative_to(data_dir)) for p in test_paths]},
        calibration=evaluate(calibration_records, calibration_scores, threshold),
        test=evaluate(test_records, test_scores, threshold),
        limitations="Only labeled curves are used. Other loading points are assumed negative; "
                     "unannotated pop-ins may count as false positives. Scores are not calibrated probabilities.")
    # Preserve the tested model: don't refit it on the held-out curves.
    # Replace only after a complete successful write, keeping the last usable
    # model if training or serialization fails.
    temporary_model = model_path.with_suffix(".joblib.tmp")
    try:
        joblib.dump(dict(version=MODEL_VERSION, feature_schema=FEATURE_SCHEMA, classifier=classifier,
                         threshold=threshold, report=report), temporary_model, compress=3)
        temporary_model.replace(model_path)
    finally:
        temporary_model.unlink(missing_ok=True)
    model_path.with_suffix(".report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


def load_model(path=None):
    # joblib uses pickle: only load models generated locally or otherwise trusted.
    if path is None:
        path = DEFAULT_MODEL if DEFAULT_MODEL.exists() else BUNDLED_MODEL
    bundle = joblib.load(path)
    if not isinstance(bundle, dict) or bundle.get("version") != MODEL_VERSION:
        raise DataError("Unsupported model. Retrain with this version of the program.")
    if not {"classifier", "threshold", "report"}.issubset(bundle):
        raise DataError("Incomplete model. Retrain the model.")
    if bundle.get("feature_schema") != FEATURE_SCHEMA or bundle["classifier"].n_features_in_ != 60:
        raise DataError("Incompatible model features. Click Train Shape Model to rebuild it.")
    return bundle


def predict(curve, bundle):
    c = loading(curve)
    probabilities = bundle["classifier"].predict_proba(features(curve))[:, 1]
    return [dict(index=i, depth_nm=float(c[i, 0]), load_uN=float(c[i, 1]),
                 time_s=float(c[i, 2]), score=float(probabilities[i]))
            for i in event_indices(probabilities, bundle["threshold"])]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    training = sub.add_parser("train")
    training.add_argument("--data", type=Path, default=DEFAULT_DATA)
    training.add_argument("--labels", type=Path, default=DEFAULT_LABELS)
    training.add_argument("--model", type=Path, default=DEFAULT_MODEL)
    inference = sub.add_parser("predict")
    inference.add_argument("file", type=Path)
    inference.add_argument("--model", type=Path, default=DEFAULT_MODEL)
    args = parser.parse_args()
    try:
        result = (train(args.data, args.labels, args.model) if args.command == "train"
                  else predict(load_curve(args.file), load_model(args.model)))
    except Exception as exc:
        parser.exit(1, f"{type(exc).__name__}: {exc}\n")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()

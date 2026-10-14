import json
from pathlib import Path

import numpy as np
import pytest
import joblib
from openpyxl import Workbook

from analysis_scripts.popin_ml import (DEFAULT_DATA, DEFAULT_LABELS, SHEET_FOLDERS, DataError,
                      evaluate, event_indices, features, load_curve,
                      match_labels, read_labels)
from analysis_scripts.popin_ml import DEFAULT_MODEL, collect_training_data, load_model, predict, train


def write_curve(path, points=100, corrupt=None, time_scale=1):
    rows = ["Instrument export", "Depth (nm)\tLoad (uN)\tTime (s)\tDepth (V)\tLoad (V)"]
    for i in range(points):
        row = f"{i * 0.1}\t{i * 2.5}\t{i * 0.01 * time_scale}\t0\t0"
        rows.append(corrupt if i == 50 and corrupt else row)
    path.write_text("\n".join(rows))
    return path


def workbook(path, pairs):
    book = Workbook()
    book.remove(book.active)
    for name in SHEET_FOLDERS:
        sheet = book.create_sheet(name)
        sheet.append(["Depth", "Force (uN)"])
        for pair in pairs:
            sheet.append(pair)
        sheet.append([None, None])
        sheet.append([1.2, 99])  # A numeric summary must not become a label.
    book.save(path)
    return path


def test_curve_reader_and_features(tmp_path):
    curve = load_curve(write_curve(tmp_path / "curve.txt"))
    assert curve.shape == (100, 3)
    assert features(curve).shape == (100, 60)
    assert np.isfinite(features(curve)).all()


def test_features_invariant_to_axis_scaling_and_offsets():
    i = np.arange(200, dtype=float)
    depth = 0.1 * i + 0.02 * np.sin(i)
    depth[90:] += 2
    curve = np.column_stack([depth, 2.5 * i + 0.1 * np.sin(i), 0.005 * i])
    transformed = curve * [3, 2, 4] + [70, 800, 10]
    np.testing.assert_allclose(features(curve), features(transformed), rtol=1e-5, atol=1e-5)


def test_old_feature_model_rejected(tmp_path):
    path = tmp_path / "old.joblib"
    joblib.dump({"version": 1}, path)
    with pytest.raises(DataError, match="Unsupported model"):
        load_model(path)


@pytest.mark.skipif(not DEFAULT_MODEL.exists(), reason="Train model first")
def test_trained_predictions_invariant_to_scaling():
    model = load_model()
    curve = load_curve(DEFAULT_DATA / "SiC-4H-OFF" / "SiC-4H 0000.txt")
    scaled = curve * [1.7, 2, 3] + [50, 300, 8]
    original_events, scaled_events = predict(curve, model), predict(scaled, model)
    assert [e["index"] for e in original_events] == [e["index"] for e in scaled_events]
    np.testing.assert_allclose([e["score"] for e in original_events],
                               [e["score"] for e in scaled_events], atol=1e-5)


@pytest.mark.parametrize("bad", ["bad\t2\t3\t0\t0", "nan\t2\t3\t0\t0", "1\t2"])
def test_bad_measurements_rejected(tmp_path, bad):
    with pytest.raises(DataError, match="Line"):
        load_curve(write_curve(tmp_path / "bad.txt", corrupt=bad))


def test_empty_short_wrong_extension_and_time(tmp_path):
    for filename, content in (("empty.txt", ""), ("wrong.csv", "1,2,3")):
        path = tmp_path / filename
        path.write_text(content)
        with pytest.raises(DataError):
            load_curve(path)
    with pytest.raises(DataError, match="50"):
        load_curve(write_curve(tmp_path / "short.txt", points=10))
    path = write_curve(tmp_path / "time.txt")
    path.write_text(path.read_text().replace("5.0\t125.0\t0.5", "5.0\t125.0\t0.49"))
    with pytest.raises(DataError, match="Time"):
        load_curve(path)


def test_wrong_units_are_rejected(tmp_path):
    path = write_curve(tmp_path / "units.txt")
    path.write_text(path.read_text().replace("Load (uN)", "Load (mN)"))
    with pytest.raises(DataError, match="microNewtons"):
        load_curve(path)


def test_excel_summary_excluded(tmp_path):
    labels = read_labels(workbook(tmp_path / "labels.xlsx", [(5, 125)]))
    assert len(labels) == 4
    assert all(label["excel_row"] == 2 for label in labels)


def test_matching_requires_both_coordinates_and_detects_ambiguity(tmp_path):
    labels = workbook(tmp_path / "labels.xlsx", [(5, 125), (5, 126)])
    for number, folder in enumerate(SHEET_FOLDERS.values(), 1):
        directory = tmp_path / folder
        directory.mkdir()
        write_curve(directory / "sample 0000.txt", time_scale=number)
    records, audit = match_labels(tmp_path, labels)
    assert len(records) == 4
    assert [row["status"] for row in audit].count("matched") == 4
    assert [row["status"] for row in audit].count("unmatched") == 4
    write_curve(tmp_path / SHEET_FOLDERS["Off"] / "duplicate 0001.txt")
    records, audit, files = collect_training_data(tmp_path, labels)
    assert len(records) == 4
    assert audit[0]["status"] == "matched"
    assert any(row["status"] == "duplicate_curve" for row in files)


def test_scan_arbitrary_names_nested_files_and_skip_invalid_unlabeled(tmp_path):
    labels = workbook(tmp_path / "labels.xlsx", [(5, 125), (5, 126)])
    folder = tmp_path / SHEET_FOLDERS["Off"] / "new batch"
    folder.mkdir(parents=True)
    write_curve(folder / "new experiment.TXT")
    (folder / "analysis.txt").write_text("not a curve")
    write_curve(folder / "no label.txt", points=50)
    records, audit, files = collect_training_data(tmp_path, labels)
    assert len(records) == 1
    assert next(iter(records)).name == "new experiment.TXT"
    assert {row["status"] for row in files} == {"matched", "invalid_format", "no_matching_annotation"}
    assert any(row["status"] == "unmatched" for row in audit)


def test_training_skips_files_and_saves_reloadable_model(tmp_path):
    from openpyxl import load_workbook
    import shutil
    labels = tmp_path / "labels.xlsx"
    shutil.copy(DEFAULT_LABELS, labels)
    book = load_workbook(labels)
    book["Off"].insert_rows(2)
    book["Off"]["A2"] = 99999
    book["Off"]["B2"] = 99999  # An unmatched annotation must not stop training.
    book.save(labels)
    data = tmp_path / "training"
    shutil.copytree(DEFAULT_DATA, data)
    (data / "invalid.txt").write_text("not a curve")
    path = tmp_path / "saved" / "model.joblib"
    report = train(data, labels, path)
    assert report["labeled_curves"] == 115
    assert report["unmatched_annotation_count"] == 1
    assert any(row["file"] == "invalid.txt" for row in report["skipped_files"])
    assert path.with_suffix(".files.csv").exists()
    curve = load_curve(data / "SiC-4H-OFF" / "SiC-4H 0000.txt")
    first = predict(curve, load_model(path))
    reopened = predict(curve, load_model(path))
    assert [event["index"] for event in first] == [event["index"] for event in reopened]
    np.testing.assert_allclose([event["score"] for event in first],
                               [event["score"] for event in reopened], atol=1e-12)
    # Failure on insufficient eligible curves preserves the saved model.
    previous = path.read_bytes()
    small = tmp_path / "too_small"
    small.mkdir()
    (small / "bad.txt").write_text("invalid")
    with pytest.raises(DataError, match="At least 20"):
        train(small, labels, path)
    assert path.read_bytes() == previous


def test_event_grouping_and_one_to_one_scoring():
    scores = np.zeros(100)
    scores[20:23] = [0.6, 0.8, 0.7]
    scores[26] = 0.9
    scores[70] = 0.8
    assert event_indices(scores, 0.5) == [26, 70]
    result = evaluate([{"indices": [25, 90]}], [scores], 0.5)
    assert (result["true_positives"], result["false_positives"], result["false_negatives"]) == (1, 1, 1)


def test_repository_annotations_and_disjoint_splits():
    records, audit = match_labels(DEFAULT_DATA, DEFAULT_LABELS)
    assert len(audit) == 115
    assert len(records) == 115
    assert all(row["status"] == "matched" for row in audit)
    report_path = Path(__file__).resolve().parents[1] / "models" / "popin_model.report.json"
    if report_path.exists():
        splits = json.loads(report_path.read_text())["split_files"]
        sets = [set(splits[key]) for key in ("train", "calibration", "test")]
        assert sum(map(len, sets)) == 115
        assert not (sets[0] & sets[1] or sets[0] & sets[2] or sets[1] & sets[2])
